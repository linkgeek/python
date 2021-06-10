import xlrd
import random
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# 读取excel
data = xlrd.open_workbook('data3.xlsx')
sheet = data.sheet_by_index(0)


class Question:
    pass

# 获取试题
def create_question():
    questions = []
    for i in range(sheet.nrows):
        if i > 1:
            obj = Question()
            obj.subject = sheet.cell(i, 1).value
            obj.type = sheet.cell(i, 2).value
            obj.option = []
            obj.option.append(sheet.cell(i, 3).value)
            obj.option.append(sheet.cell(i, 4).value)
            obj.option.append(sheet.cell(i, 5).value)
            obj.option.append(sheet.cell(i, 6).value)
            obj.score = sheet.cell(i, 7)
            questions.append(obj)
    # 随机排序
    random.shuffle(questions)
    return questions


# 生成word试卷
def createPaper(filename, papername, list):

    doc = Document()
    # 页眉
    section = doc.sections[0]
    header = section.header
    p1 = header.paragraphs[0]
    p1.text = papername
    footer = section.footer
    p2 = footer.paragraphs[0]
    p2.text = "内部试卷，禁止泄露"

    # 页头
    title = doc.add_heading(papername, level=1)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p3 = doc.add_paragraph()
    p3.add_run('姓名：')
    p3.add_run('所属部门：')
    p3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # 试题信息
    for item in list:
        subject = doc.add_paragraph(style='List Number')
        run = subject.add_run(item.subject)
        run.bold = True
        subject.add_run('【%s】'% str(item.score))
        for idx, option in enumerate(item):
            doc.add_paragraph(('ABCD')[idx] + str(option))
    doc.save('../data/' + filename)


data = create_question()
createPaper('001', '2020年第一季度内部考试', data)